"""
cal_cert_importer.py

Parses a Keysight "Certificate of Calibration" + "Calibration Report" PDF
(as combined by Keysight into one file) and appends the new data point
into the matching resistor block of a trending workbook, then extends
the linear-regression / prediction-interval formulas to the new row.

Usage:
    python cal_cert_importer.py <workbook.xlsx> <cert.pdf> [--yes]

Requires an "Asset_Index" sheet in the workbook with columns:
    Serial Number | Sheet Name | Nominal Value | Date Col | Since Col |
    Accum Col | Error Col | Drift Col | Measured Col | MU Col |
    Pred Col | Lower Col | Upper Col

Design notes / limitations (read before trusting this in production):
  - Page 1/2 of the cert (Certificate of Calibration) has a real text
    layer -> Serial Number and Date of Calibration are parsed directly.
  - The "Calibration Report" page (Nominal/Measured/MU table) is often
    a SCANNED IMAGE with no text layer, so it's OCR'd with tesseract.
    OCR occasionally misreads the Ohm symbol as "Q" or "O", the SI unit
    prefix (k/M/G/m/u), and inserts stray spaces inside numbers. The
    measured value's scale is cross-checked against the *known* nominal
    already recorded in Asset_Index (not the OCR'd nominal), since a
    misread unit prefix (e.g. m vs M) is a 1e9x error and far more
    dangerous than the Ohm-symbol confusion.
  - Because OCR is never 100% guaranteed, the script ALWAYS prints what
    it extracted (plus any correction warnings) and asks for confirmation
    before writing anything. Pass --yes only once you trust it for a
    given cert layout.
  - Rows are assumed to be entered in chronological order (matches
    current practice). The script appends after the last row of the
    matched block; it does not re-sort.
"""

import argparse
import logging
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import pdfplumber
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.formula import ArrayFormula

ASSET_INDEX_SHEET = "Asset_Index"
LOG_PATH = Path(__file__).parent / "cal_cert_importer.log"

logger = logging.getLogger("cal_cert_importer")


def _setup_logging():
    logger.setLevel(logging.INFO)
    if logger.handlers:
        return  # already configured (e.g. re-imported in tests)
    fh = logging.FileHandler(LOG_PATH, encoding="utf-8")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(fh)

# SI prefixes as they're likely to appear on these certs. Case matters a lot
# here (m = milli = 1e-3, M = mega = 1e6 - a 1e9 difference), which is
# exactly the kind of OCR mistake this script must never trust blindly.
UNIT_MULTIPLIERS = {
    "p": 1e-12, "n": 1e-9, "u": 1e-6, "\u00b5": 1e-6, "m": 1e-3,
    "": 1.0, "k": 1e3, "K": 1e3, "M": 1e6, "G": 1e9,
}

# Matches "<number, may have internal spaces/dot> <optional SI prefix> <Ohm-like symbol>"
# e.g. "10.000 002 kQ" -> ("10.000 002", "k"); ".100 000 761 Q" -> (".100 000 761", "")
_VALUE_UNIT_RE = re.compile(r"([\d.][\d.\s]*)\s*([pnu\u00b5mkKMG]?)\s*[\u03a9QO]")


# ---------------------------------------------------------------- PDF parsing

def extract_serial_and_date(pdf_path: str, tesseract_cmd: str = None) -> dict:
    """Extract serial, calibration date, and certificate number from text or scanned headers.

    Keysight certificate formats vary: combined certificates often use
    "Serial Number" / "Date of Calibration", while Fluke 732A/C reports may
    be scanned and use "S/N" / "DATE". OCR is used only for pages without a
    usable text layer.
    """
    chunks = []
    with pdfplumber.open(pdf_path) as pdf:
        if not pdf.pages:
            raise ValueError("The selected PDF contains no pages.")
        for page in pdf.pages[:2]:
            text = page.extract_text() or ""
            if not text.strip():
                try:
                    text = _ocr_page(page, tesseract_cmd, psm=3)
                except RuntimeError as exc:
                    logger.warning("Header OCR failed for %s: %s", pdf_path, exc)
                    text = ""
            chunks.append(text)
    header_text = "\n".join(chunks)

    serial_patterns = (
        r"Serial\s+Number\s*[:#]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)",
        r"\bS\s*/\s*N\s*[:#]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)",
        r"\bSER(?:IAL)?\.?\s*NO\.?\s*[:#]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)",
    )
    serial = None
    for pattern in serial_patterns:
        match = re.search(pattern, header_text, re.I)
        if match:
            serial = match.group(1).strip(" .,:;")
            break

    date_patterns = (
        r"Date\s+of\s+Calibration\s*:?\s*(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
        r"\bDATE\s*:?\s*(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})",
    )
    cal_date = None
    for pattern in date_patterns:
        match = re.search(pattern, header_text, re.I)
        if not match:
            continue
        raw_date = re.sub(r"\s+", " ", match.group(1).strip())
        for fmt in ("%d %b %Y", "%d %B %Y"):
            try:
                cal_date = datetime.strptime(raw_date, fmt)
                break
            except ValueError:
                continue
        if cal_date:
            break

    cert_match = re.search(r"\b(?:CERT(?:IFICATE)?\s*(?:NO\.?|NUMBER)?\s*[:#]?\s*:?\s*)(WO[- ]?\d{6,})", header_text, re.I)
    cert_number = re.sub(r"\s+", "", cert_match.group(1)).replace(" ", "-") if cert_match else None
    return {"serial": serial, "cal_date": cal_date, "cert_number": cert_number}


def extract_measurement(pdf_path: str, known_nominal: float = None, tesseract_cmd: str = None) -> dict:
    """Extract the legacy resistor result only; other categories need point-level data."""
    diagnostics = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                try:
                    text = _ocr_page(page, tesseract_cmd)
                except RuntimeError as exc:
                    raise RuntimeError(f"Could not OCR PDF page {page_number}: {exc}") from exc
            if "MEASURED VALUE" in text.upper():
                measured_value, mu_relative, warnings = _parse_measurement_line(text, known_nominal)
                if measured_value is not None and mu_relative is not None:
                    return {"measured_value": measured_value, "mu_relative": mu_relative, "warnings": warnings}
                diagnostics.append(f"Page {page_number} contains 'MEASURED VALUE' but the resistor value/MU row could not be parsed.")
            elif text.strip():
                diagnostics.append(f"Page {page_number}: no 'MEASURED VALUE' table header found.")
    if not diagnostics:
        diagnostics.append("No readable text or measurement table was found in the PDF.")
    return {"measured_value": None, "mu_relative": None, "warnings": diagnostics}


_TESSERACT_FALLBACK_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    "/usr/bin/tesseract",
    "/usr/local/bin/tesseract",
]


def _find_tesseract() -> str:
    found = shutil.which("tesseract")
    if found:
        return found
    for candidate in _TESSERACT_FALLBACK_PATHS:
        if Path(candidate).exists():
            return candidate
    raise RuntimeError(
        "Tesseract OCR was not found. Install Tesseract (Windows installer: "
        "https://github.com/UB-Mannheim/tesseract/wiki) or configure its executable path."
    )


def _ocr_page(page, tesseract_cmd: str = None, psm: int = 6) -> str:
    """OCR one page using a unique temporary image; never terminates the host app."""
    cmd = tesseract_cmd or _find_tesseract()
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(prefix="cal_cert_", suffix=".png", delete=False) as tmp:
            temp_path = Path(tmp.name)
        page.to_image(resolution=300 if psm == 3 else 600).save(str(temp_path))
        try:
            result = subprocess.run(
                [cmd, str(temp_path), "stdout", "--psm", str(psm)],
                capture_output=True, text=True, timeout=90, check=False,
            )
        except FileNotFoundError as exc:
            raise RuntimeError(f"Tesseract executable '{cmd}' could not be started.") from exc
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("Tesseract OCR timed out after 90 seconds.") from exc
        if result.returncode != 0:
            detail = (result.stderr or "").strip()
            raise RuntimeError(f"Tesseract failed (exit {result.returncode}): {detail or 'no error details'}")
        return result.stdout
    finally:
        if temp_path is not None:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                logger.warning("Could not remove temporary OCR image: %s", temp_path)


def _clean_number(raw: str) -> float:
    s = re.sub(r"[^\d.]", "", raw)
    if s.count(".") > 1:  # OCR occasionally duplicates a stray dot
        first = s.index(".")
        s = s[: first + 1] + s[first + 1:].replace(".", "")
    s = s.rstrip(".")  # trailing stray dot only - a leading dot is a real decimal point
    return float(s) if s and s != "." else 0.0


def _parse_measurement_line(text: str, known_nominal: float = None) -> tuple:
    """Returns (measured_value_ohms, mu_relative, warnings)."""
    candidates = [l for l in text.splitlines() if len(re.findall(r"[\u03a9QO]", l)) >= 2]

    for line in candidates:
        found = list(_VALUE_UNIT_RE.finditer(line))
        if len(found) < 2:
            continue

        warnings = []
        nominal_raw = (found[0].group(1), found[0].group(2))
        measured_raw = (found[1].group(1), found[1].group(2))
        measured_num = _clean_number(measured_raw[0])
        if measured_num == 0:
            continue

        if len(found) >= 3:
            mu_raw = (found[2].group(1), found[2].group(2))
            mu_num = _clean_number(mu_raw[0])
            mu_prefix = mu_raw[1]
        else:
            # MU's unit got mangled badly enough that the full value+prefix+symbol
            # pattern didn't match at all (e.g. a stray OCR letter wedged in front
            # of the prefix). Fall back to just grabbing the next number after the
            # measured value's match and assume the near-universal uOhm/Ohm unit -
            # the plausibility clamp below still guards against nonsense.
            rest = line[found[1].end():]
            mu_num_match = re.search(r"\d+\.\d+|\d+", rest)
            if not mu_num_match:
                continue
            mu_num = float(mu_num_match.group())
            mu_prefix = None
            warnings.append(
                "The MU unit/prefix was not fully recognized by OCR. The numeric MU "
                "was interpreted as µΩ/Ω for preview only; verify it against the PDF before import."
            )

        # --- resolve the measured value's scale ---
        if known_nominal is not None:
            ocr_multiplier = UNIT_MULTIPLIERS.get(measured_raw[1], 1.0)
            candidate = measured_num * ocr_multiplier
            if 0.5 * known_nominal < candidate < 1.5 * known_nominal:
                measured_value = candidate
            else:
                fixed = None
                for prefix, mult in UNIT_MULTIPLIERS.items():
                    trial = measured_num * mult
                    if 0.5 * known_nominal < trial < 1.5 * known_nominal:
                        fixed = (trial, prefix)
                        break
                if fixed is None:
                    continue  # nothing plausible on this line - try the next candidate
                measured_value, corrected_prefix = fixed
                warnings.append(
                    f"OCR read the measured-value unit prefix as '{measured_raw[1]}', which "
                    f"didn't match the known nominal ({known_nominal} ohm). Assumed "
                    f"'{corrected_prefix}' instead based on the known nominal - VERIFY "
                    f"this against the certificate before confirming."
                )
        else:
            measured_value = measured_num * UNIT_MULTIPLIERS.get(measured_raw[1], 1.0)
            nominal_guess = _clean_number(nominal_raw[0]) * UNIT_MULTIPLIERS.get(nominal_raw[1], 1.0)
            if nominal_guess > 0 and not (0.5 * nominal_guess < measured_value < 1.5 * nominal_guess):
                continue
            warnings.append(
                "No known nominal was available to validate the unit prefix - double-check "
                "the measured value's magnitude carefully."
            )

        # --- resolve MU: always a relative ratio (e.g. uOhm/Ohm), so its
        # prefix must land in a plausible ppm-scale band rather than being
        # trusted verbatim ---
        mu_relative = mu_num * UNIT_MULTIPLIERS.get(mu_prefix, 1e-6)
        if not (1e-9 < mu_relative < 1e-3):
            # Do not silently accept an implausible uncertainty unit. Keep the
            # historical µΩ/Ω assumption for compatibility, but make it explicit.
            warnings.append(
                "The extracted MU magnitude/unit was implausible. It was interpreted "
                "as µΩ/Ω for preview only; verify the uncertainty against the PDF."
            )
            mu_relative = mu_num * 1e-6

        return measured_value, mu_relative, warnings

    return None, None, []



# ---------------------------------------------------------------- Multi-point certificate extraction

def _number_without_group_spaces(raw: str) -> float:
    """Convert OCR grouped numbers such as '10.000 083 883' to a float."""
    cleaned = re.sub(r"(?<=\d)\s+(?=\d{3}(?:\s|$))", "", raw.strip())
    cleaned = re.sub(r"\s+", "", cleaned)
    return float(cleaned)


def _normalise_frequency(value: str) -> str:
    s = re.sub(r"\s+", "", value).lower()
    s = s.replace("khz", "khz").replace("mhz", "mhz")
    if s.endswith("khz"):
        number = s[:-3]
        return f"{number} kHz"
    if s.endswith("mhz"):
        number = s[:-3]
        return f"{number} MHz"
    return s


def extract_multi_point_measurements(
    pdf_path: str,
    category: str,
    test_points: list,
    tesseract_cmd: str = None,
) -> dict:
    """Extract known multi-point Keysight/Fluke certificate layouts.

    Supported layouts validated against supplied examples:
      * Fluke 732A/C DCV report: nominal output, measured voltage, uncertainty (nV)
      * Keysight AC divider report: frequency, ratio, uncertainty (ppm)

    The function returns readings keyed to configured database test-point IDs.
    It never returns a partial result as a complete import.
    """
    page_texts = []
    diagnostics = []
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                # Scanned measurement pages have no text layer. OCR with a layout
                # mode that tends to preserve table rows/columns.
                if not text.strip():
                    try:
                        if category == "ac_divider":
                            # The validated ET divider report places its six-row
                            # ratio table in the lower-middle of the scanned page.
                            # Cropping improves row ordering; full-page OCR tends
                            # to read each column independently on this layout.
                            crop = page.crop((35, 245, page.width - 35, min(page.height, 555)))
                            text = _ocr_page(crop, tesseract_cmd, psm=6)
                        else:
                            text = _ocr_page(page, tesseract_cmd, psm=3)
                    except RuntimeError as exc:
                        diagnostics.append(f"Page {page_number} OCR failed: {exc}")
                if text.strip():
                    page_texts.append((page_number, text))
    except Exception as exc:
        return {"ok": False, "readings": [], "warnings": [f"Could not read PDF: {exc}"]}

    all_text = "\n".join(text for _, text in page_texts)
    extracted = {}

    if category == "dcv_reference":
        # Example rows: '10 10.000 083 883 35' and
        # '1 0.999 982 438 24.4'. Preserve the grouped fractional digits.
        row_re = re.compile(
            r"^\s*(?:(10|1)\s+)?((?:\d+\.\d+|\.\d+)(?:\s+\d{3}){1,3})\s+([\d]+(?:\.\d+)?)\s*$"
        )
        for line in all_text.splitlines():
            line = line.replace("|", " ").strip()
            match = row_re.match(line)
            if not match:
                continue
            try:
                measured = _number_without_group_spaces(match.group(2))
                uncertainty_nv = float(match.group(3))
            except ValueError:
                continue
            nominal = float(match.group(1)) if match.group(1) else (10.0 if measured > 5.0 else 1.0)
            if not (8.0 <= measured <= 12.0 or 0.5 <= measured <= 1.5):
                continue
            key = "10v" if nominal == 10 else "1v"
            if key in extracted:
                diagnostics.append(f"Duplicate {nominal:g} V output row found; refusing ambiguous extraction.")
                continue
            extracted[key] = {"measured_value": measured, "uncertainty": uncertainty_nv}

        expected_keys = {}
        for tp in test_points:
            name = str(tp.get("point_name", tp.get("name", ""))).lower().replace(" ", "")
            nominal = tp.get("nominal_value")
            if "10v" in name or (nominal is not None and abs(float(nominal) - 10.0) < 1e-9):
                expected_keys[tp.get("id")] = "10v"
            elif "1v" in name or (nominal is not None and abs(float(nominal) - 1.0) < 1e-9):
                expected_keys[tp.get("id")] = "1v"

    elif category == "ac_divider":
        # Row-oriented OCR is important here: each frequency, ratio and ppm
        # uncertainty must come from the same line to avoid column-shift errors.
        row_re = re.compile(
            r"^\s*(1\s*MHz|500\s*kHz|300\s*kHz|100\s*kHz|20\s*kHz|1\s*kHz)\s+"
            r"([\d]+(?:\.\d+)?)\s+([\d]+(?:\.\d*)?)\s*ppm\b",
            re.I,
        )
        for line in all_text.splitlines():
            match = row_re.match(line.strip())
            if not match:
                continue
            freq = _normalise_frequency(match.group(1))
            try:
                measured = float(match.group(2))
                uncertainty_ppm = float(match.group(3).rstrip("."))
            except ValueError:
                continue
            if freq in extracted:
                diagnostics.append(f"Duplicate {freq} ratio row found; refusing ambiguous extraction.")
                continue
            extracted[freq] = {"measured_value": measured, "uncertainty": uncertainty_ppm}

        expected_keys = {
            tp.get("id"): _normalise_frequency(str(tp.get("point_name", tp.get("name", "")))
                                                .replace(" ", ""))
            for tp in test_points
        }
    else:
        return {
            "ok": False, "readings": [],
            "warnings": [f"Automatic multi-point extraction is not yet supported for category '{category}'."],
        }

    readings = []
    missing = []
    for tp in test_points:
        tp_id = tp.get("id")
        key = expected_keys.get(tp_id)
        value = extracted.get(key)
        if value is None:
            missing.append(str(tp.get("point_name", tp.get("name", tp_id))))
            continue
        readings.append({
            "test_point_id": tp_id,
            "measured_value": value["measured_value"],
            "uncertainty": value["uncertainty"],
        })

    if missing:
        diagnostics.append("Could not confidently extract configured point(s): " + ", ".join(missing))
    if not test_points:
        diagnostics.append("No configured test points were supplied for matching.")
    ok = bool(test_points) and not missing and len(readings) == len(test_points)
    warnings = [
        "OCR-extracted values must be checked against the original certificate before saving."
    ]
    warnings.extend(diagnostics)
    return {"ok": ok, "readings": readings, "warnings": warnings,
            "uncertainty_unit": {"dcv_reference": "nV", "ac_divider": "ppm"}.get(category)}


# ---------------------------------------------------------------- Workbook side

def load_asset_index(wb) -> dict:
    if ASSET_INDEX_SHEET not in wb.sheetnames:
        sys.exit(f"Workbook has no '{ASSET_INDEX_SHEET}' sheet.")
    ws = wb[ASSET_INDEX_SHEET]
    index = {}
    headers = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[col["Serial Number"]]:
            continue
        index[str(row[col["Serial Number"]])] = {
            "sheet": row[col["Sheet Name"]],
            "nominal": row[col["Nominal Value"]],
            "cols": {
                "date": row[col["Date Col"]],
                "since": row[col["Since Col"]],
                "accum": row[col["Accum Col"]],
                "error": row[col["Error Col"]],
                "drift": row[col["Drift Col"]],
                "measured": row[col["Measured Col"]],
                "mu": row[col["MU Col"]],
                "pred": row[col["Pred Col"]],
                "lower": row[col["Lower Col"]],
                "upper": row[col["Upper Col"]],
            },
        }
    return index


def _last_data_row(ws, date_col: str) -> int:
    row = 3
    last = 2
    while ws[f"{date_col}{row}"].value is not None:
        last = row
        row += 1
    return last


def _copy_formula(ws, col: str, src_row: int, dst_row: int):
    src_cell = ws[f"{col}{src_row}"]
    dst_cell = ws[f"{col}{dst_row}"]
    val = src_cell.value
    if val is None:
        return
    if isinstance(val, ArrayFormula):
        new_text = Translator(val.text, origin=src_cell.coordinate).translate_formula(dst_cell.coordinate)
        dst_cell.value = ArrayFormula(dst_cell.coordinate, new_text)
    elif isinstance(val, str) and val.startswith("="):
        dst_cell.value = Translator(val, origin=src_cell.coordinate).translate_formula(dst_cell.coordinate)
    else:
        dst_cell.value = val


def append_record(workbook_path: str, record: dict, asset: dict):
    wb = openpyxl.load_workbook(workbook_path, keep_vba=workbook_path.endswith(".xlsm"))
    ws = wb[asset["sheet"]]
    cols = asset["cols"]
    last_row = _last_data_row(ws, cols["date"])
    new_row = last_row + 1

    existing_date = ws[f"{cols['date']}{last_row}"].value
    if existing_date and record["cal_date"] and existing_date >= record["cal_date"]:
        print(f"WARNING: new date {record['cal_date'].date()} is not after last "
              f"recorded date {existing_date.date()} for this asset. Appending anyway "
              f"at row {new_row} - reorder manually if needed.")

    for key in ("since", "accum", "error", "drift", "pred", "lower", "upper"):
        col = cols[key]
        if col:
            _copy_formula(ws, col, last_row, new_row)

    ws[f"{cols['date']}{new_row}"] = record["cal_date"]
    ws[f"{cols['measured']}{new_row}"] = record["measured_value"]
    ws[f"{cols['mu']}{new_row}"] = round(record["mu_relative"] * asset["nominal"], 12)

    wb.save(workbook_path)
    print(f"Wrote row {new_row} on sheet '{asset['sheet']}' of {workbook_path}")


# ---------------------------------------------------------------- CLI

def main():
    _setup_logging()
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("cert_pdf")
    ap.add_argument("--yes", action="store_true", help="skip confirmation prompt")
    ap.add_argument("--tesseract-exe", default=None,
                     help=r'path to tesseract.exe if not in the default install location, e.g. "C:\path\to\tesseract.exe"')
    ap.add_argument("--db", default=None,
                     help="optional path to a Phase-2 SQLite database (db.py) to also write this "
                          "calibration into, in addition to the Excel workbook")
    args = ap.parse_args()

    logger.info("=== run start: workbook=%s cert=%s ===", args.workbook, args.cert_pdf)

    id_fields = extract_serial_and_date(args.cert_pdf)
    print(f"Serial Number : {id_fields['serial']}")
    print(f"Cal Date      : {id_fields['cal_date'].date() if id_fields['cal_date'] else None}")

    if id_fields["serial"] is None or id_fields["cal_date"] is None:
        logger.error("Could not read Serial Number / Date of Calibration from %s", args.cert_pdf)
        sys.exit("Could not read Serial Number / Date of Calibration - aborting.")

    wb = openpyxl.load_workbook(args.workbook, keep_vba=args.workbook.endswith(".xlsm"))
    index = load_asset_index(wb)
    asset = index.get(str(id_fields["serial"]))
    if asset is None:
        logger.error("Serial %s not found in %s (%s)", id_fields["serial"], ASSET_INDEX_SHEET, args.workbook)
        sys.exit(f"Serial number {id_fields['serial']} not found in {ASSET_INDEX_SHEET}. "
                  f"Add it there first (with its column block + nominal value).")

    measurement = extract_measurement(args.cert_pdf, known_nominal=asset["nominal"], tesseract_cmd=args.tesseract_exe)
    print(f"Measured Value: {measurement['measured_value']}")
    print(f"MU (relative) : {measurement['mu_relative']}")
    for w in measurement["warnings"]:
        print(f"  ! {w}")
        logger.warning("%s | serial=%s cert=%s", w, id_fields["serial"], args.cert_pdf)

    if measurement["measured_value"] is None or measurement["mu_relative"] is None:
        logger.error("Failed to parse measurement for serial=%s cert=%s", id_fields["serial"], args.cert_pdf)
        sys.exit("Could not parse Measured Value / MU - aborting. Check the OCR output above.")

    record = {**id_fields, **measurement}

    if not args.yes or measurement["warnings"]:
        answer = input("Write this into the workbook? [y/N] ").strip().lower()
        if answer != "y":
            logger.info("User declined to write serial=%s cert=%s", id_fields["serial"], args.cert_pdf)
            sys.exit("Aborted - nothing written.")

    append_record(args.workbook, record, asset)
    logger.info("Wrote to workbook: serial=%s date=%s measured=%s mu=%s",
                id_fields["serial"], id_fields["cal_date"].date(),
                measurement["measured_value"], measurement["mu_relative"])

    if args.db:
        from db import connect, upsert_asset, insert_calibration
        conn = connect(args.db)
        asset_id = upsert_asset(
            conn, serial=id_fields["serial"], model=None, asset_number=None,
            nominal_ohms=asset["nominal"], source_workbook=args.workbook, source_sheet=asset["sheet"],
        )
        cert_number_m = re.search(r"(WO-\d+)", Path(args.cert_pdf).stem)
        inserted = insert_calibration(
            conn, asset_id, cal_date=id_fields["cal_date"].date().isoformat(),
            measured_value=measurement["measured_value"], mu_relative=measurement["mu_relative"],
            source="pdf_import", cert_number=cert_number_m.group(1) if cert_number_m else None,
            cert_pdf_path=str(Path(args.cert_pdf).resolve()),
        )
        conn.commit()
        conn.close()
        msg = "Also wrote to %s" if inserted else "Already present in %s (same asset+date) - skipped"
        print(msg % args.db)
        logger.info(msg.replace("%s", args.db))


if __name__ == "__main__":
    main()
