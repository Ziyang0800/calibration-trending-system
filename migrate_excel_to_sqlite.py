"""
migrate_excel_to_sqlite.py

Migrates historical calibration records into calibration.db:
  1. Legacy Decade Resistors (10-column blocks)
  2. DC Voltage Reference (Fluke 732A / 732C) with Barometric Pressure
  3. AC Voltage Dividers (ET 10mV / ET 100mV, 6 frequencies)
  4. Waveform Generators (Keysight 33622A, 15 points with Adjustment detection)

Usage:
    python migrate_excel_to_sqlite.py <file1.xlsx> [<file2.xlsx> ...] [--db calibration.db]
"""

import argparse
import re
import sys
from pathlib import Path
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from db import (
    connect, upsert_asset, insert_calibration,
    upsert_test_point, insert_calibration_reading
)

BLOCK_WIDTH = 10
HEADER_ROW2 = ["Date", "Since Last Cal [Days]", "Accumulated [Days]", "Error [PPM]",
               "Drifted [PPM]", "Measured Value", "MU (K =2)", "Prediction", "Lower", "Upper"]

UNIT_MULTIPLIERS = {
    "p": 1e-12, "n": 1e-9, "u": 1e-6, "\u00b5": 1e-6, "m": 1e-3,
    "": 1.0, "k": 1e3, "K": 1e3, "M": 1e6, "G": 1e9,
}
_NOMINAL_NAME_RE = re.compile(r"([\d.]+)\s*([pnu\u00b5mkKMG]?)\s*Ohm", re.IGNORECASE)


def parse_nominal_from_sheet_name(sheet_name: str) -> float:
    m = _NOMINAL_NAME_RE.search(sheet_name)
    if not m:
        return None
    value, prefix = m.groups()
    return float(value) * UNIT_MULTIPLIERS.get(prefix, 1.0)


def _fmt_date(val) -> str:
    """Parses Excel dates (datetime, date, or strings like 9/3/2019 or 2019-09-03)."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if "/" in s:
        parts = s.split("/")
        if len(parts) == 3:
            m, d, y = parts
            y = f"20{y}" if len(y) == 2 else y
            return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    return s[:10] if len(s) >= 10 else s


def _safe_float(val):
    if val is None or val == "":
        return None
    try:
        return float(str(val).replace(",", "").strip())
    except ValueError:
        return None


def _read_meta_block(ws):
    meta = {}
    for r in range(1, 6):
        lbl = ws.cell(row=r, column=1).value
        val = ws.cell(row=r, column=2).value
        if lbl:
            meta[str(lbl).strip()] = str(val).strip() if val is not None else ""
    return meta


# -----------------------------------------------------------------------------
# 1. New Templates (DCV Reference, AC Divider, Waveform Generator)
# -----------------------------------------------------------------------------
def migrate_sheet_template(ws, path: str, conn) -> int:
    meta = _read_meta_block(ws)
    category_raw = meta.get("Category", "")
    model = meta.get("Model")
    serial = meta.get("Serial Number")
    equip_id = meta.get("Equipment ID")
    set_id = meta.get("Set ID")

    if not serial:
        return 0

    total_new = 0

    # A. DC Voltage Reference (Fluke 732A / 732C) - 10V and 1V Support
    if "DC Voltage" in category_raw or "732" in str(model):
        asset_id = upsert_asset(
            conn, serial=serial, model=model, asset_number=equip_id,
            nominal_ohms=10.0, source_workbook=Path(path).name, source_sheet=ws.title,
            category="dcv_reference", unit="V", set_id=set_id
        )
        # Setup both 10V and 1V test points
        tp_10v = upsert_test_point(conn, asset_id, "10V Output", "DC Voltage", 10.0, "V", 1)
        tp_1v = upsert_test_point(conn, asset_id, "1V Output", "DC Voltage", 1.0, "V", 2)

        # Check if sheet has the new 1V columns (Columns D and E)
        has_1v_cols = "1V" in str(ws.cell(row=7, column=4).value or "")

        row = 8
        while ws.cell(row=row, column=1).value is not None:
            cal_date = _fmt_date(ws.cell(row=row, column=1).value)
            meas_10v = _safe_float(ws.cell(row=row, column=2).value)
            uncert_10v = _safe_float(ws.cell(row=row, column=3).value)

            if has_1v_cols:
                meas_1v = _safe_float(ws.cell(row=row, column=4).value)
                uncert_1v = _safe_float(ws.cell(row=row, column=5).value)
                location = ws.cell(row=row, column=6).value
                pressure = _safe_float(ws.cell(row=row, column=7).value)
                cert_no = ws.cell(row=row, column=8).value
                comment = ws.cell(row=row, column=9).value
            else:
                meas_1v = None
                uncert_1v = None
                cert_no = ws.cell(row=row, column=4).value
                pressure = _safe_float(ws.cell(row=row, column=5).value)
                comment_val = ws.cell(row=row, column=6).value
                # If comment was a facility name, store in location
                if comment_val and str(comment_val).strip() in ("Loveland", "Roseville"):
                    location = str(comment_val).strip()
                    comment = None
                else:
                    location = None
                    comment = comment_val

            if meas_10v is not None and cal_date:
                mu_rel_10v = (uncert_10v * 1e-9 / 10.0) if uncert_10v is not None else 0.0
                cal_id = insert_calibration(
                    conn, asset_id, cal_date, meas_10v, mu_rel_10v,
                    source="excel_migration", cert_number=str(cert_no) if cert_no else None,
                    pressure_hpa=pressure, reviewer_comment=str(comment) if comment else None,
                    location=str(location) if location else None
                )
                if cal_id:
                    # Save 10V reading
                    insert_calibration_reading(conn, cal_id, tp_10v, meas_10v, uncertainty=uncert_10v)
                    # Save 1V reading if present
                    if meas_1v is not None:
                        insert_calibration_reading(conn, cal_id, tp_1v, meas_1v, uncertainty=uncert_1v)
                    total_new += 1
            row += 1

    # B. AC Voltage Divider (ET 10mV / ET 100mV)
    elif "Divider" in category_raw or "ET " in str(model) or "ETAC" in str(model):
        nom_ratio = 100.0 if "10mV" in str(model) or "10MV" in str(model) else 10.0
        asset_id = upsert_asset(
            conn, serial=serial, model=model, asset_number=equip_id,
            nominal_ohms=nom_ratio, source_workbook=Path(path).name, source_sheet=ws.title,
            category="ac_divider", unit="Ratio", set_id=set_id
        )

        freqs = [("1 kHz", 2, 3), ("20 kHz", 4, 5), ("100 kHz", 6, 7),
                 ("300 kHz", 8, 9), ("500 kHz", 10, 11), ("1 MHz", 12, 13)]
        tp_ids = {}
        for idx, (f_name, _, _) in enumerate(freqs, start=1):
            tp_ids[f_name] = upsert_test_point(conn, asset_id, f_name, "Ratio Characterization", nom_ratio, "Ratio", idx)

        row = 8
        while ws.cell(row=row, column=1).value is not None:
            cal_date = _fmt_date(ws.cell(row=row, column=1).value)
            primary_ratio = _safe_float(ws.cell(row=row, column=2).value)
            primary_uncert = _safe_float(ws.cell(row=row, column=3).value)
            cert_no = ws.cell(row=row, column=14).value
            comment = ws.cell(row=row, column=15).value

            if primary_ratio is not None and cal_date:
                mu_rel = (primary_uncert * 1e-6) if primary_uncert is not None else 0.0
                cal_id = insert_calibration(
                    conn, asset_id, cal_date, primary_ratio, mu_rel,
                    source="excel_migration", cert_number=str(cert_no) if cert_no else None,
                    reviewer_comment=str(comment) if comment else None
                )
                if cal_id:
                    for f_name, r_col, u_col in freqs:
                        r_val = _safe_float(ws.cell(row=row, column=r_col).value)
                        u_val = _safe_float(ws.cell(row=row, column=u_col).value)
                        if r_val is not None:
                            insert_calibration_reading(conn, cal_id, tp_ids[f_name], r_val, uncertainty=u_val)
                    total_new += 1
            row += 1

    # C. Waveform Generator (Keysight 33622A)
    elif "Waveform" in category_raw or "33622" in str(model):
        asset_id = upsert_asset(
            conn, serial=serial, model=model, asset_number=equip_id,
            nominal_ohms=0.0, source_workbook=Path(path).name, source_sheet=ws.title,
            category="waveform_generator", unit="dB", set_id=set_id
        )

        cf_headers = [
            "cf 10mV @ 4MHz (dB)", "cf 100mV @ 4MHz (dB)", "cf 100mV @ 8MHz (dB)",
            "cf 100mV @ 10MHz (dB)", "cf 1V @ 4MHz (dB)", "cf 1V @ 8MHz (dB)",
            "cf 1V @ 10MHz (dB)", "cf 3V @ 2MHz (dB)", "cf 3V @ 4MHz (dB)",
            "cf 3V @ 8MHz (dB)", "cf 3V @ 10MHz (dB)"
        ]
        flat_headers = [
            "Flatness Error 2MHz (dB)", "Flatness Error 4MHz (dB)",
            "Flatness Error 8MHz (dB)", "Flatness Error 10MHz (dB)"
        ]

        tp_ids = {}
        for idx, h in enumerate(cf_headers, start=1):
            tp_ids[idx] = upsert_test_point(conn, asset_id, h, "Calibration Factor (cf)", 0.0, "dB", idx)
        for idx, h in enumerate(flat_headers, start=12):
            tp_ids[idx] = upsert_test_point(conn, asset_id, h, "Flatness Error", 0.0, "dB", idx)

        row = 8
        while ws.cell(row=row, column=1).value is not None:
            cal_date = _fmt_date(ws.cell(row=row, column=1).value)
            primary_cf = _safe_float(ws.cell(row=row, column=2).value)
            cert_no = ws.cell(row=row, column=17).value
            comment = ws.cell(row=row, column=18).value

            # Automatic Adjustment Detection:
            is_adj = bool(comment and "adjust" in str(comment).lower())

            if primary_cf is not None and cal_date:
                cal_id = insert_calibration(
                    conn, asset_id, cal_date, primary_cf, 0.0,
                    source="excel_migration", cert_number=str(cert_no) if cert_no else None,
                    is_adjusted=is_adj, reviewer_comment=str(comment) if comment else None
                )
                if cal_id:
                    for col_idx in range(1, 16):
                        val = _safe_float(ws.cell(row=row, column=col_idx + 1).value)
                        if val is not None:
                            insert_calibration_reading(conn, cal_id, tp_ids[col_idx], val)
                    total_new += 1
            row += 1

    print(f"  [{category_raw or model}] {serial} ({ws.title}): {total_new} calibration(s) migrated.")
    return total_new


# -----------------------------------------------------------------------------
# 2. Legacy Decade Resistors
# -----------------------------------------------------------------------------
def find_blocks(ws):
    max_col = ws.max_column
    for col_idx in range(1, max_col):
        c1 = ws.cell(row=2, column=col_idx).value
        c2 = ws.cell(row=2, column=col_idx + 1).value
        if c1 == "Date" and c2 == "Since Last Cal [Days]":
            cols = [get_column_letter(col_idx + i) for i in range(BLOCK_WIDTH)]
            actual = [ws[f"{c}2"].value for c in cols]
            if actual != HEADER_ROW2:
                continue
            model = ws[f"{cols[1]}1"].value
            asset_number = ws[f"{cols[2]}1"].value
            serial = ws[f"{cols[3]}1"].value
            yield {
                "cols": dict(zip(["date", "since", "accum", "error", "drift", "measured", "mu", "pred", "lower", "upper"], cols)),
                "model": str(model) if model is not None else None,
                "asset_number": str(asset_number) if asset_number is not None else None,
                "serial": str(serial) if serial is not None else None,
            }


def migrate_legacy_workbook(path: str, db_path: str) -> int:
    wb_cached = openpyxl.load_workbook(path, data_only=True)
    conn = connect(db_path)
    total_new = 0

    for ws in wb_cached.worksheets:
        if ws.title == "Asset_Index":
            continue
        nominal = parse_nominal_from_sheet_name(ws.title)
        blocks = list(find_blocks(ws))
        if not blocks or nominal is None:
            continue

        for block in blocks:
            if block["serial"] is None:
                continue
            asset_id = upsert_asset(
                conn, serial=block["serial"], model=block["model"],
                asset_number=block["asset_number"], nominal_ohms=nominal,
                source_workbook=Path(path).name, source_sheet=ws.title,
                category="standard_resistor", unit="Ω"
            )
            tp_id = upsert_test_point(conn, asset_id, f"{nominal:g} Ω Nominal", "Resistance", nominal, "Ω", 1)

            date_col = block["cols"]["date"]
            measured_col = block["cols"]["measured"]
            mu_col = block["cols"]["mu"]
            row = 3
            block_new = 0
            while ws[f"{date_col}{row}"].value is not None:
                date_val = ws[f"{date_col}{row}"].value
                measured_val = _safe_float(ws[f"{measured_col}{row}"].value)
                mu_val = _safe_float(ws[f"{mu_col}{row}"].value)

                if measured_val is not None and mu_val is not None:
                    mu_relative = mu_val / nominal
                    date_str = _fmt_date(date_val)
                    cal_id = insert_calibration(conn, asset_id, date_str, measured_val, mu_relative, source="excel_migration")
                    if cal_id:
                        insert_calibration_reading(conn, cal_id, tp_id, measured_val, uncertainty=mu_val)
                        block_new += 1
                row += 1
            print(f"  [Standard Resistor] {ws.title} / {block['serial']} ({block['asset_number']}): {block_new} row(s) migrated.")
            total_new += block_new

    conn.commit()
    conn.close()
    return total_new


def migrate_file(path: str, db_path: str) -> int:
    wb = openpyxl.load_workbook(path, data_only=True)
    is_template = any(ws.cell(row=1, column=1).value == "Category" for ws in wb.worksheets)

    if is_template:
        conn = connect(db_path)
        total = 0
        for ws in wb.worksheets:
            if ws.cell(row=1, column=1).value == "Category":
                total += migrate_sheet_template(ws, path, conn)
        conn.commit()
        conn.close()
        return total
    else:
        wb.close()
        return migrate_legacy_workbook(path, db_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbooks", nargs="+")
    ap.add_argument("--db", default="calibration.db")
    args = ap.parse_args()

    grand_total = 0
    for path in args.workbooks:
        print(f"Migrating {path} ...")
        grand_total += migrate_file(path, args.db)

    print(f"\nDone. {grand_total} calibration records migrated to {args.db}.")


if __name__ == "__main__":
    main()